from datetime import date
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "SRS_Sistema_Gestion_Muebleria.docx"


SPANISH_REPLACEMENTS = {
    "Especificacion": "Especificación",
    "Gestion": "Gestión",
    "Muebleria": "Mueblería",
    "administracion": "administración",
    "Administracion": "Administración",
    "facturacion": "facturación",
    "autenticacion": "autenticación",
    "Autenticacion": "Autenticación",
    "accion": "acción",
    "Accion": "Acción",
    "sesion": "sesión",
    "Sesion": "Sesión",
    "paginas": "páginas",
    "Paginas": "Páginas",
    "codigo": "código",
    "modulo": "módulo",
    "modulos": "módulos",
    "Modulo": "Módulo",
    "descripcion": "descripción",
    "Descripcion": "Descripción",
    "Tecnologia": "Tecnología",
    "Aplicacion": "Aplicación",
    "aplicacion": "aplicación",
    "analisis": "análisis",
    "Analisis": "Análisis",
    "Version": "Versión",
    "proposito": "propósito",
    "Proposito": "Propósito",
    "operacion": "operación",
    "Operacion": "Operación",
    "navegacion": "navegación",
    "Navegacion": "Navegación",
    "validacion": "validación",
    "Validacion": "Validación",
    "numero": "número",
    "Numero": "Número",
    "clasificacion": "clasificación",
    "Clasificacion": "Clasificación",
    "lineas": "líneas",
    "Lineas": "Líneas",
    "pequena": "pequeña",
    "pequeno": "pequeño",
    "basico": "básico",
    "basica": "básica",
    "tambien": "también",
    "historico": "histórico",
    "Historico": "Histórico",
    "Categorias": "Categorías",
    "Categorias": "Categorías",
    "categorias": "categorías",
    "categoria": "categoría",
    "Categoria": "Categoría",
    "Termino": "Término",
    "Definicion": "Definición",
    "Area": "Área",
    "Situacion": "Situación",
    "Recomendacion": "Recomendación",
    "Proteccion": "Protección",
    "Seleccion": "Selección",
    "calculos": "cálculos",
    "calculo": "cálculo",
    "Calculo": "Cálculo",
    "actualizacion": "actualización",
    "Actualizacion": "Actualización",
    "Generacion": "Generación",
    "vacio": "vacío",
    "boton": "botón",
    "impresion": "impresión",
    "cedula": "cédula",
    "telefono": "teléfono",
    "direccion": "dirección",
    "informacion": "información",
    "Informacion": "Información",
    "ultimas": "últimas",
    "Ultimas": "Últimas",
    "mas": "más",
    "disponible": "disponible",
    "ejecucion": "ejecución",
    "configuracion": "configuración",
    "Configuracion": "Configuración",
    "segun": "según",
    "Segun": "Según",
    "interactua": "interactúa",
    "dolares": "dólares",
    "exposicion": "exposición",
    "academica": "académica",
    "evolucion": "evolución",
}


def apply_spanish_replacements_to_text(text):
    for old, new in SPANISH_REPLACEMENTS.items():
        text = re.sub(rf"(?<!\w){re.escape(old)}(?!\w)", new, text)
    return text


def polish_spanish_text(doc):
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.text = apply_spanish_replacements_to_text(run.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.text = apply_spanish_replacements_to_text(run.text)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def style_table(table, widths):
    set_table_width(table, widths)
    table.style = "Table Grid"
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
            if row_idx == 0:
                set_cell_shading(cell, "F2F4F7")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(31, 77, 120)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for idx, text in enumerate(headers):
        hdr[idx].text = text
    for row_data in rows:
        row = table.add_row().cells
        for idx, text in enumerate(row_data):
            row[idx].text = str(text)
    style_table(table, widths)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_footer(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "SRS - Sistema de Gestión para Mueblería"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(89, 89, 89)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Documento generado a partir del proyecto Flask actual")
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(89, 89, 89)


def add_cover(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Especificación de Requerimientos del Sistema (SRS)")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Sistema de Gestión para Mueblería")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(46, 116, 181)

    doc.add_paragraph()
    metadata = [
        ("Proyecto", "Sistema Flask de administración, inventario, clientes, ventas y facturación"),
        ("Tecnología base", "Python Flask, Jinja2, MySQL, HTML y CSS"),
        ("Arquitectura observada", "Aplicación web modular con Blueprints, modelos de acceso a datos y plantillas"),
        ("Versión del documento", "1.0"),
        ("Fecha", date.today().strftime("%d/%m/%Y")),
        ("Fuente", "Análisis del código actual del proyecto"),
    ]
    add_table(doc, ["Campo", "Detalle"], metadata, [2700, 6660])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "Este informe describe los requerimientos actuales y esperados del sistema, "
        "tomando como base los módulos implementados en la aplicación."
    ).italic = True
    doc.add_page_break()


def main():
    doc = Document()
    configure_document(doc)
    add_footer(doc)
    add_cover(doc)

    doc.add_heading("1. Introduccion", level=1)
    doc.add_paragraph(
        "El presente documento SRS define los requerimientos funcionales y no funcionales "
        "del Sistema de Gestion para Muebleria. El sistema esta construido como una aplicacion "
        "web en Flask y permite administrar usuarios autenticados, productos, categorias, clientes, "
        "ventas, facturas e indicadores operativos del negocio."
    )
    doc.add_paragraph(
        "El objetivo del documento es servir como referencia formal para desarrollo, pruebas, "
        "mantenimiento y futuras mejoras del proyecto."
    )

    doc.add_heading("1.1 Proposito", level=2)
    add_bullets(
        doc,
        [
            "Documentar los requisitos del sistema actualmente trabajado.",
            "Definir el alcance funcional de los modulos existentes.",
            "Establecer criterios de aceptacion para validar el comportamiento del sistema.",
            "Identificar restricciones tecnicas, reglas de negocio y riesgos relevantes.",
        ],
    )

    doc.add_heading("1.2 Alcance del Sistema", level=2)
    add_bullets(
        doc,
        [
            "Ingreso seguro mediante pantalla de inicio de sesion.",
            "Dashboard con ventas del mes, facturas emitidas, productos, clientes, productos mas vendidos y alertas de stock bajo.",
            "Gestion de categorias de productos.",
            "Gestion de productos con precio, material, stock y categoria.",
            "Gestion de clientes con cedula, telefono, correo y direccion.",
            "Registro de ventas por cliente, con carrito de productos, calculo de subtotal, IVA y total.",
            "Generacion de numero de factura, detalle de venta, actualizacion de inventario e impresion de factura.",
        ],
    )

    doc.add_heading("1.3 Definiciones y Abreviaturas", level=2)
    add_table(
        doc,
        ["Termino", "Definicion"],
        [
            ("SRS", "Software Requirements Specification o Especificacion de Requerimientos del Sistema."),
            ("CRUD", "Operaciones de crear, consultar, actualizar y eliminar registros."),
            ("IVA", "Impuesto aplicado a la venta; en el sistema se calcula al 15%."),
            ("Dashboard", "Panel de control con indicadores y resumenes operativos."),
            ("Blueprint", "Modulo de Flask usado para separar rutas por funcionalidad."),
        ],
        [2200, 7160],
    )

    doc.add_heading("2. Descripcion General", level=1)
    doc.add_heading("2.1 Perspectiva del Producto", level=2)
    doc.add_paragraph(
        "El sistema es una aplicacion web administrativa para una muebleria. Opera sobre una base "
        "de datos MySQL llamada proyecto_flask y separa responsabilidades en controladores, modelos, "
        "plantillas y archivos de estilo. La navegacion principal se concentra en Dashboard, Categorias, "
        "Productos, Clientes y Ventas."
    )

    doc.add_heading("2.2 Actores", level=2)
    doc.add_paragraph(
        "De acuerdo con el proyecto actual, el sistema no implementa roles diferenciados. "
        "Las pantallas internas se habilitan para cualquier usuario autenticado; por tanto, "
        "el actor operativo se interpreta como un administrador o vendedor que usa las mismas "
        "credenciales y permisos dentro del sistema."
    )
    add_table(
        doc,
        ["Actor", "Descripcion", "Interacciones principales"],
        [
            (
                "Administrador/Vendedor (usuario autenticado)",
                "Persona con credenciales registradas en la tabla usuarios. En el sistema actual representa al usuario operativo autorizado, sin distinción técnica de roles.",
                "Accede al dashboard; administra categorias, productos y clientes; registra ventas; consulta detalles y genera facturas.",
            ),
        ],
        [1900, 3400, 4060],
    )
    doc.add_paragraph(
        "Nota: el cliente no se considera actor del sistema porque no inicia sesion ni interactua "
        "directamente con la aplicacion. Se maneja como entidad del dominio, asociada a ventas, "
        "facturas y datos de contacto."
    )

    doc.add_heading("2.3 Entorno Operativo", level=2)
    add_table(
        doc,
        ["Componente", "Descripcion"],
        [
            ("Backend", "Python con Flask, Blueprints y manejo de sesiones."),
            ("Frontend", "Plantillas HTML con Jinja2 y estilos CSS por modulo."),
            ("Base de datos", "MySQL mediante flask_mysqldb."),
            ("Servidor", "Ejecucion local o servidor compatible con Flask."),
            ("Navegador", "Cliente web moderno compatible con HTML5, CSS y JavaScript basico."),
        ],
        [2400, 6960],
    )

    doc.add_heading("2.4 Restricciones", level=2)
    add_bullets(
        doc,
        [
            "El sistema depende de una instancia MySQL disponible en localhost.",
            "La configuracion actual usa la base de datos proyecto_flask.",
            "La autenticacion implementada valida usuario y password contra la tabla usuarios.",
            "Las rutas protegidas requieren una sesion activa.",
            "La numeracion de facturas se basa en el ultimo registro de ventas.",
            "El proyecto no evidencia roles diferenciados; funciona con usuario autenticado unico.",
        ],
    )

    doc.add_heading("3. Requerimientos Funcionales", level=1)
    functional_rows = [
        ("RF-01", "Autenticacion", "El sistema debe permitir que un usuario ingrese con usuario y contraseña.", "Alta"),
        ("RF-02", "Sesion", "El sistema debe impedir el acceso a modulos internos si no existe una sesion activa.", "Alta"),
        ("RF-03", "Cierre de sesion", "El sistema debe permitir cerrar sesion y limpiar los datos almacenados en sesion.", "Alta"),
        ("RF-04", "Dashboard", "El sistema debe mostrar ventas del mes, facturas emitidas, total de productos y total de clientes.", "Media"),
        ("RF-05", "Indicadores comerciales", "El sistema debe listar productos mas vendidos, alertas de stock bajo y ultimas ventas.", "Media"),
        ("RF-06", "Categorias", "El sistema debe permitir crear, listar, editar y eliminar categorias.", "Alta"),
        ("RF-07", "Productos", "El sistema debe permitir crear, listar, editar y eliminar productos con nombre, descripcion, material, precio, stock y categoria.", "Alta"),
        ("RF-08", "Clientes", "El sistema debe permitir crear, listar, editar, eliminar y buscar clientes por nombre o cedula.", "Alta"),
        ("RF-09", "Venta en carrito", "El sistema debe permitir seleccionar cliente, producto y cantidad para construir el detalle de venta.", "Alta"),
        ("RF-10", "Validacion de stock", "El sistema debe evitar agregar al carrito cantidades superiores al stock disponible.", "Alta"),
        ("RF-11", "Calculo de importes", "El sistema debe calcular subtotal, IVA del 15% y total de la venta.", "Alta"),
        ("RF-12", "Registro de venta", "El sistema debe registrar venta, detalle de venta, usuario responsable y cliente asociado.", "Alta"),
        ("RF-13", "Actualizacion de inventario", "El sistema debe descontar del stock las cantidades vendidas al registrar la venta.", "Alta"),
        ("RF-14", "Factura", "El sistema debe generar una factura con numero, fecha, datos del cliente, detalle, subtotal, IVA y total.", "Alta"),
        ("RF-15", "Detalle historico", "El sistema debe permitir consultar el historial de ventas y el detalle de cada factura.", "Media"),
    ]
    add_table(doc, ["ID", "Modulo", "Requerimiento", "Prioridad"], functional_rows, [900, 1700, 5560, 1200])

    doc.add_heading("4. Requerimientos No Funcionales", level=1)
    non_functional_rows = [
        ("RNF-01", "Usabilidad", "Las pantallas deben presentar formularios claros, botones de accion y navegacion superior consistente."),
        ("RNF-02", "Seguridad", "Las paginas internas deben validar sesion antes de mostrar informacion del negocio."),
        ("RNF-03", "Integridad", "Las operaciones de venta deben mantener coherencia entre ventas, detalle_ventas y productos."),
        ("RNF-04", "Rendimiento", "Las consultas principales deben responder en tiempos adecuados para operacion local o de pequeña empresa."),
        ("RNF-05", "Mantenibilidad", "El codigo debe conservar separacion por controladores, modelos, plantillas y configuracion."),
        ("RNF-06", "Portabilidad", "El sistema debe ejecutarse en un entorno Python compatible con Flask y MySQL."),
        ("RNF-07", "Disponibilidad", "El sistema debe estar disponible durante la jornada operativa de la muebleria."),
        ("RNF-08", "Trazabilidad", "Cada venta debe quedar asociada a un numero de factura, fecha, cliente y usuario."),
    ]
    add_table(doc, ["ID", "Atributo", "Descripcion"], non_functional_rows, [1000, 1800, 6560])

    doc.add_heading("5. Modelo de Datos", level=1)
    doc.add_paragraph(
        "A partir de los modelos y consultas del proyecto se identifican las siguientes entidades principales:"
    )
    data_rows = [
        ("usuarios", "id_usuario, nombre, usuario, password", "Autenticacion y responsable de ventas."),
        ("clientes", "id_cliente, nombre, cedula, telefono, correo, direccion", "Registro de compradores."),
        ("categorias", "id_categoria, nombre, descripcion", "Clasificacion de productos."),
        ("productos", "id_producto, nombre, descripcion, material, precio, stock, id_categoria", "Inventario disponible para venta."),
        ("ventas", "id_venta, numero_factura, fecha, subtotal, iva, total, id_cliente, id_usuario", "Cabecera de cada venta registrada."),
        ("detalle_ventas", "id_venta, id_producto, cantidad, precio_unitario, subtotal", "Lineas de productos vendidos."),
    ]
    add_table(doc, ["Entidad", "Campos observados", "Uso principal"], data_rows, [1700, 4600, 3060])

    doc.add_heading("6. Reglas de Negocio", level=1)
    add_numbered(
        doc,
        [
            "Todo acceso a modulos internos requiere que el usuario haya iniciado sesion.",
            "Una venta debe estar asociada a un cliente seleccionado.",
            "Una venta no puede registrarse si el carrito esta vacio.",
            "La cantidad solicitada de un producto no debe superar el stock actual.",
            "El IVA se calcula como el 15% del subtotal.",
            "El total se calcula como subtotal mas IVA.",
            "El numero de factura debe seguir el formato FAC-000001 y aumentar secuencialmente.",
            "Al confirmar la venta, el stock de cada producto vendido debe disminuir segun la cantidad registrada.",
            "Los productos con stock menor o igual a 5 deben considerarse de bajo stock para alertas.",
        ],
    )

    doc.add_heading("7. Interfaces del Sistema", level=1)
    interface_rows = [
        ("Login", "/login", "Ingreso de credenciales y mensaje de error si no son validas."),
        ("Dashboard", "/dashboard", "Resumen de indicadores, ultimas ventas, stock bajo y productos mas vendidos."),
        ("Categorias", "/categorias", "Formulario de nueva categoria y tabla de registros existentes."),
        ("Productos", "/productos", "Formulario de nuevo producto, selector de categoria y tabla de inventario."),
        ("Clientes", "/clientes", "Formulario de cliente y listado con datos de contacto."),
        ("Ventas", "/ventas", "Seleccion de cliente/producto, carrito, calculos y registro de venta."),
        ("Detalle de venta", "/ventas/detalle/<id>", "Informacion del cliente, venta y productos facturados."),
        ("Factura", "/ventas/factura/<id>", "Vista imprimible con datos de factura y boton de impresion."),
    ]
    add_table(doc, ["Pantalla", "Ruta", "Descripcion"], interface_rows, [1900, 2300, 5160])

    doc.add_heading("8. Flujos Principales", level=1)
    doc.add_heading("8.1 Inicio de sesion", level=2)
    add_numbered(
        doc,
        [
            "El usuario abre la pantalla de login.",
            "Ingresa usuario y contraseña.",
            "El sistema consulta la tabla usuarios.",
            "Si las credenciales son validas, se crea la sesion y se redirige al dashboard.",
            "Si las credenciales no son validas, se muestra un mensaje de error.",
        ],
    )

    doc.add_heading("8.2 Registro de venta", level=2)
    add_numbered(
        doc,
        [
            "El usuario ingresa al modulo Ventas.",
            "Selecciona un cliente, un producto y una cantidad.",
            "El sistema valida existencia de producto y stock disponible.",
            "El producto se agrega al carrito con precio y subtotal.",
            "El sistema calcula subtotal, IVA y total.",
            "El usuario registra la venta.",
            "El sistema genera numero de factura, guarda cabecera y detalle, descuenta inventario y limpia el carrito.",
        ],
    )

    doc.add_heading("9. Criterios de Aceptacion", level=1)
    acceptance_rows = [
        ("CA-01", "Login", "Con credenciales validas, el usuario debe acceder al dashboard."),
        ("CA-02", "Proteccion de rutas", "Al abrir una ruta interna sin sesion, el sistema debe redirigir al login."),
        ("CA-03", "Productos", "Al crear un producto, debe aparecer en el listado con precio, stock y categoria."),
        ("CA-04", "Clientes", "Al actualizar un cliente, el listado debe reflejar los nuevos datos."),
        ("CA-05", "Stock", "Si la cantidad supera el stock, el sistema debe impedir agregar el producto a la venta."),
        ("CA-06", "Venta", "Al registrar venta valida, debe crearse cabecera, detalle y total correcto."),
        ("CA-07", "Inventario", "Tras registrar venta, el stock debe disminuir segun las unidades vendidas."),
        ("CA-08", "Factura", "La factura debe mostrar datos del cliente, detalle, subtotal, IVA y total."),
    ]
    add_table(doc, ["ID", "Area", "Criterio"], acceptance_rows, [1000, 1700, 6660])

    doc.add_heading("10. Riesgos y Recomendaciones", level=1)
    risk_rows = [
        ("Seguridad de contraseñas", "El codigo compara password directamente contra la base de datos.", "Usar hash seguro de contraseñas y migrar usuarios existentes."),
        ("Configuracion sensible", "La clave de MySQL y secret_key estan en codigo fuente.", "Mover credenciales a variables de entorno o archivo de configuracion no versionado."),
        ("Eliminaciones", "Las eliminaciones se realizan por rutas GET.", "Cambiar a POST/DELETE y agregar proteccion CSRF."),
        ("Integridad de venta", "La venta y cada detalle hacen commits separados.", "Usar transacciones para asegurar rollback si falla una parte del registro."),
        ("Roles", "No se observan permisos por perfil.", "Agregar roles si el negocio requiere separar administrador y vendedor."),
        ("Validaciones", "Algunos campos dependen de validacion HTML y conversion basica en backend.", "Validar rangos, formatos y datos obligatorios tambien en servidor."),
    ]
    add_table(doc, ["Riesgo", "Situacion actual", "Recomendacion"], risk_rows, [2200, 3500, 3660])

    doc.add_heading("11. Supuestos y Dependencias", level=1)
    add_bullets(
        doc,
        [
            "La base de datos contiene las tablas esperadas por los modelos del proyecto.",
            "Los usuarios son creados previamente en la tabla usuarios.",
            "Los navegadores clientes permiten ejecutar JavaScript basico para comportamiento de formularios.",
            "El sistema esta orientado a una muebleria con volumen operativo pequeno o medio.",
            "Los precios y totales se manejan en dolares, de acuerdo con las vistas actuales.",
        ],
    )

    doc.add_heading("12. Conclusion", level=1)
    doc.add_paragraph(
        "El proyecto actual cubre los procesos esenciales de una muebleria: autenticacion, control de inventario, "
        "registro de clientes, ventas, facturacion y monitoreo de indicadores. El SRS permite consolidar estos "
        "comportamientos como base documental para pruebas, exposicion academica, mejoras de seguridad y evolucion "
        "funcional del sistema."
    )

    polish_spanish_text(doc)
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
